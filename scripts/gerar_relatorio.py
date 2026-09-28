"""Gera README.md e o relatório Word a partir dos mesmos blocos de conteúdo e de data/processed/resultados.json.

Uso: python scripts/gerar_relatorio.py [caminho_docx]
"""
import json
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PROC, EVID = ROOT / "data" / "processed", ROOT / "evidencias"
DOCX_OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "Relatorio_MVP_Industria_Juros_Lucas_Iwano.docx"
R = json.loads((PROC / "resultados.json").read_text(encoding="utf-8"))
LOG = json.loads((PROC / "log_etl.json").read_text(encoding="utf-8"))
MAN = json.loads((ROOT / "data" / "raw" / "_manifesto_coleta.json").read_text(encoding="utf-8"))
CAT = pd.read_csv(ROOT / "catalogo" / "catalogo_dados.csv")
P1 = pd.read_csv(PROC / "tab_P1_recuperacao.csv", dtype={"codigo": str})
P5 = pd.read_csv(PROC / "tab_P5_cenario.csv", dtype={"codigo": str})
NUVEM = json.loads((EVID / "nuvem_status.json").read_text(encoding="utf-8")) if (EVID / "nuvem_status.json").exists() else {"executado": False}
REPO = "https://github.com/iwanolucas/MVP-Industria-Juros"


def f(x, d=1):
    return f"{x:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def pct(x, d=1):
    return f(x, d) + "%"


def nome(c):
    return P1.set_index("codigo").atividade[c]


B = []  # blocos: (tipo, ...)
h1 = lambda t: B.append(("h1", t))
h2 = lambda t: B.append(("h2", t))
p = lambda t: B.append(("p", t))
bl = lambda items: B.append(("ul", items))
tb = lambda head, rows, widths=None: B.append(("table", head, rows, widths))
im = lambda arq, leg: B.append(("img", arq, leg))

P1r, P2r, P3r, P4r, P5r = R["P1"], R["P2"], R["P3"], R["P4"], R["P5"]

# =====================================================================================================
h1("1. Objetivo")
h2("1.1 Problema")
p("O juro real brasileiro está em nível elevado (9,3% a.a. em agosto/2026, pelos dados deste trabalho). Um gestor industrial, um banco de "
  "desenvolvimento ou um investidor que precise decidir onde expandir capacidade, provisionar risco ou renegociar prazos "
  "depara-se com duas perguntas: (i) a indústria brasileira se recuperou de fato do choque de 2020, ou a média esconde setores "
  "em depressão? e (ii) o quanto os juros e o câmbio explicam o ritmo de produção de cada setor, de modo que um corte de juros "
  "pudesse orientar a priorização de setores?")
p("O MVP constrói um pipeline de dados na nuvem que integra duas fontes públicas distintas, a Pesquisa Industrial Mensal – Produção "
  "Física (PIM-PF/IBGE), com 27 atividades industriais, e as séries macroeconômicas do Banco Central (Selic, IPCA e dólar), e usa o "
  "resultado para testar hipóteses explícitas sobre a relação entre política monetária, câmbio e produção industrial setorial.")
p("O objetivo foi definido antes da busca e da análise dos dados. A base foi escolhida depois, por ser aberta, agregada (sem dados "
  "pessoais) e capaz de responder às perguntas abaixo.")
h2("1.2 Hipóteses")
bl([
    "H1 — Juros. Juro real alto reduz o crescimento da produção industrial, com defasagem de alguns meses, e o efeito é mais forte em setores de bens de capital e de consumo durável, dependentes de crédito.",
    "H2 — Câmbio. A desvalorização do real está associada a maior produção nos setores exportadores e de substituição de importações, depois de controlado o efeito dos juros.",
    "H3 — Heterogeneidade. A recuperação pós-pandemia foi desigual: o agregado voltou ao nível de 2019, mas há divisões industriais que ainda estão abaixo dele.",
])
h2("1.3 Perguntas de negócio")
tb(["#", "Pergunta", "Decisão a que se liga"], [
    ["P1", "Quais atividades industriais ainda estão abaixo do nível de produção de 2019 e quais o superaram? Em quantos meses cada uma voltou ao patamar pré-pandemia?", "Onde há capacidade ociosa estrutural e onde há demanda aquecida."],
    ["P2", "Quais atividades têm crescimento anual mais sensível ao juro real e com que defasagem (0 a 12 meses)? A relação é estatisticamente distinguível do acaso?", "Quando esperar o efeito de uma mudança da Selic e em que setores."],
    ["P3", "Depois de controlar o juro real, a variação cambial de 12 meses está associada à produção de cada atividade? Em quais?", "Exposição cambial (insumos importados vs. competitividade)."],
    ["P4", "Quais atividades têm maior volatilidade do crescimento anual e quão maior é o risco relativo entre elas?", "Provisão de risco e dimensionamento de capital de giro."],
    ["P5", "Se o juro real caísse 3 p.p. (da faixa atual de 9,3% a.a. para perto da média histórica de 5,5%), qual variação de crescimento anual o modelo projeta para a indústria e para cada divisão, e com que incerteza?", "Priorização de setores em um cenário de flexibilização monetária."],
])
p("Nenhuma pergunta foi removida ao final: as que não puderam ser respondidas de forma conclusiva estão discutidas na Seção 5.6 e na Autoavaliação.")

# =====================================================================================================
h1("2. Coleta")
h2("2.1 Fontes")
tb(["Arquivo bruto", "Origem", "Registros", "Tamanho", "Formato original"],
   [[m["arquivo"], m["url"].replace("https://", "")[:70] + ("…" if len(m["url"]) > 78 else ""), f(m["registros"], 0), f(m["bytes"] / 1024, 0) + " KB", "JSON"] for m in MAN])
p(f"Coleta realizada em {MAN[0]['coletado_em_utc'][:10].split('-')[2]}/{MAN[0]['coletado_em_utc'][:10].split('-')[1]}/{MAN[0]['coletado_em_utc'][:4]} "
  "(UTC) pelo script scripts/coleta.py, que grava também data/raw/_manifesto_coleta.json com URL, data, volume, formato e hash SHA-256 de cada arquivo. "
  "As APIs são públicas e não exigem chave. O script repete a requisição até 5 vezes, porque a API do BCB devolveu corpo vazio de forma intermitente durante o desenvolvimento.")
p("Não houve web scraping: as duas fontes oferecem API oficial. A tabela 8888 da SIDRA traz o número-índice de produção física com base 2022 = 100, "
  "sem ajuste sazonal, para o Brasil, de janeiro/2002 a julho/2026 (295 meses × 27 atividades). As quatro séries do SGS são mensais: "
  "Selic acumulada no mês anualizada (4189), IPCA mensal (433), IPCA acumulado em 12 meses (13522) e dólar médio de venda (3698).")
h2("2.2 Persistência na nuvem")
if NUVEM.get("executado"):
    p(NUVEM.get("texto_persistencia", ""))
else:
    p("Os arquivos brutos ficam versionados no repositório (data/raw). A persistência na plataforma de nuvem (Databricks, Volume do Unity Catalog e "
      "tabelas Delta) é feita pelos notebooks 01 e 02, e o status da execução consta na Seção 4.3.")

# =====================================================================================================
h1("3. Modelagem")
h2("3.1 Modelo de dados")
p("Adotou-se um esquema dimensional do tipo constelação de fatos, com duas tabelas fato que compartilham a dimensão tempo:")
bl([
    "fato_producao: uma linha por mês e atividade (7.965 linhas): índice de produção, variação anual e indicador de ausência.",
    "fato_macro: uma linha por mês (297 linhas): Selic, IPCA, dólar e as métricas derivadas (juro real e variação cambial de 12 meses).",
    "dim_tempo: 297 meses, de janeiro/2002 a setembro/2026, com ano, trimestre e marcador de pandemia.",
    "dim_atividade: 27 atividades com código CNAE, nível hierárquico (geral, seção, divisão) e nome abreviado.",
])
p("A escolha se justifica por três razões. As duas fontes têm o mesmo grão temporal (mês), o que torna o tempo a dimensão natural de integração. "
  "As métricas macro pertencem ao mês, não à atividade, e repeti-las em cada linha de produção multiplicaria por 27 o volume e o risco de inconsistência. "
  "Por fim, o esquema permite responder às perguntas com joins simples de SQL.")
im(ROOT / "catalogo" / "diagrama_estrela.png", "Modelo dimensional do MVP (catalogo/diagrama_estrela.png).")
h2("3.2 Catálogo de dados")
p("Os domínios abaixo foram lidos da própria base carregada, não declarados de memória. O catálogo completo, com descrição e linhagem de cada atributo, "
  "está em catalogo/catalogo_dados.md e .csv.")
tb(["Tabela", "Atributo", "Tipo", "Domínio observado", "Nulos", "Linhagem (resumo)"],
   [[r.tabela, r.atributo, r.tipo, r.dominio_observado, str(r.nulos_observados), r.linhagem] for r in CAT.itertuples()])
p("Sobre a obrigatoriedade: os únicos nulos com significado próprio são (a) o índice das divisões 3.18 (Impressão) e 3.33 (Manutenção), que o IBGE só "
  "publica a partir de 2012, (b) os 12 primeiros meses de variação anual e cambial, que dependem de um ano de histórico, e (c) os campos do mês corrente do BCB "
  "(setembro/2026), ainda sem divulgação completa. Nenhum nulo foi imputado.")

# =====================================================================================================
h1("4. Carga (ETL)")
h2("4.1 Arquitetura do pipeline")
p("O pipeline segue a arquitetura em camadas (medallion) e é executado por três notebooks Databricks, cada um um invólucro fino sobre um script Python versionado. "
  "O mesmo código roda localmente (SQLite) e na nuvem (tabelas Delta), o que permitiu validar a lógica antes de subir à plataforma.")
tb(["Camada", "Conteúdo", "Onde", "Código"], [
    ["Bronze", "JSON bruto das APIs, sem alteração", "Volume raw (Unity Catalog) / data/raw", "notebooks/01_coleta_bronze.py → scripts/coleta.py"],
    ["Silver", "Tabelas limpas e tipadas (datas, códigos, nulos)", "Memória do ETL", "scripts/etl.py (R1 a R6)"],
    ["Gold", "Esquema estrela em Delta", "workspace.industria_juros.*", "notebooks/02_etl_silver_gold.py → scripts/etl.py (R7 a R9)"],
    ["Consumo", "Consultas SQL, testes e gráficos", "Notebook 03", "notebooks/03_analise.py → scripts/analise.py"],
])
h2("4.2 Regras de negócio da transformação")
tb(["Regra", "Descrição", "Efeito medido"], [
    ["R1", "Mês em português (\"janeiro 2002\") convertido em data (1º dia do mês) e chave AAAAMM.", "295 meses PIM-PF e 297 meses BCB alinhados."],
    ["R2", "Código CNAE separado do nome da atividade.", "27 atividades: 1 geral, 2 seções e 24 divisões."],
    ["R3", "\"-\" (sem informação) vira NULL; nunca imputado.", f"{LOG['pim_valores_ausentes_R3']} valores nulos (2 divisões × 120 meses)."],
    ["R4", "Datas do BCB (dd/mm/aaaa) normalizadas para o 1º dia do mês; valores em ponto flutuante.", "Junção exata por mês."],
    ["R5", "Duplicatas pela chave natural (mês, atividade) removidas, mantendo a primeira.", f"{LOG['pim_duplicatas_removidas_R5']} duplicatas."],
    ["R6", "Variação anual = índice ÷ índice de 12 meses antes − 1, por atividade.", "564 nulos estruturais (12 meses iniciais e ausências)."],
    ["R7", "Juro real ex-post = (1 + Selic a.a.) ÷ (1 + IPCA 12m) − 1.", "Faixa observada de −4,4% a 13,0% a.a."],
    ["R8", "Variação cambial 12m = dólar médio ÷ dólar médio de 12 meses antes − 1.", "Positivo = real desvalorizado."],
    ["R9", "Meses só do BCB (sem PIM-PF) permanecem na fato_macro com flag_tem_producao = 0.", f"{LOG['macro_meses_sem_producao_R9']} meses (ago. e set./2026)."],
])
p("Foram verificadas as integridades referencial (toda chave das fatos existe nas dimensões) e de unicidade (nenhuma duplicata na chave composta), com asserções no ETL: o script falha se alguma delas for violada.")
h2("4.3 Execução na nuvem e persistência")
if NUVEM.get("executado"):
    for par in NUVEM.get("paragrafos", []):
        p(par)
    for arq, leg in NUVEM.get("figuras", []):
        im(ROOT / "evidencias" / arq, leg)
else:
    p("Status: os notebooks Databricks (notebooks/01, 02 e 03) foram escritos e a lógica que eles executam foi validada localmente, com o mesmo código e as mesmas "
      "consultas SQL (o resultado da consulta P1 em SQL bate com o do pandas, diferença 0,0). **A execução dentro do Databricks ainda não foi realizada nesta versão** "
      "porque exige login da própria conta do aluno na plataforma, e não há evidência de nuvem (capturas de tela) neste relatório. Os passos para executar estão no README: "
      "criar a conta gratuita Databricks Free Edition, criar uma Git folder apontando para o repositório e executar os três notebooks em ordem. "
      "Ao final, o notebook 02 grava as quatro tabelas Delta e o comando DESCRIBE HISTORY comprova a persistência.")

# =====================================================================================================
h1("5. Análise")
h2("5.1 Qualidade dos dados")
p("Todos os atributos foram verificados nas seis dimensões usuais. Os números abaixo são gerados pelo script de análise (data/processed/resultados.json).")
tb(["Dimensão", "Verificação", "Resultado", "Tratamento"], [
    ["Completude", "Nulos por atributo",
     f"Índice: {R['q_indice_ausentes_total']} nulos, todos nas divisões 3.18 e 3.33 entre jan/2002 e dez/2011. IPCA e dólar: 1 nulo cada (set./2026).",
     "Não imputados. As séries são usadas só a partir de 2013, período em que todas as 27 estão completas."],
    ["Unicidade", "Duplicatas em (mês, atividade)", f"{R['q_duplicatas_chave']} duplicatas; {R['q_lacunas_temporais']} lacunas de mês em qualquer série.", "Nenhum."],
    ["Consistência", "Variação anual recalculada; IPCA 12m oficial vs. composição do IPCA mensal; indústria geral vs. seções",
     f"Diferença máxima da variação anual: {R['q_consist_yoy_max_diff']:.0e} p.p.; IPCA 12m: diferença média de {f(R['q_ipca12m_vs_composto_medio_pp'], 4)} p.p. (máxima {f(R['q_ipca12m_vs_composto_max_diff_pp'], 3)} p.p., arredondamento); a indústria geral fica entre as seções em {f(R['q_geral_entre_secoes_pct'], 0)}% dos meses.",
     "Nenhum. Fontes coerentes entre si."],
    ["Conformidade", "Aderência ao domínio",
     f"Índice > 0 em 100% das linhas (mín. {f(R['q_indice_min_max_global'][0], 1)}; máx. {f(R['q_indice_min_max_global'][1], 1)}). Selic {f(R['q_macro_dominios']['selic_aa'][0], 1)}–{f(R['q_macro_dominios']['selic_aa'][1], 1)}% a.a.; dólar R$ {f(R['q_macro_dominios']['dolar_medio'][0], 2)}–{f(R['q_macro_dominios']['dolar_medio'][1], 2)}.",
     "Nenhum. Todos plausíveis."],
    ["Acurácia", "Outliers da variação anual (|z robusto| > 5)",
     f"{R['q_outliers_yoy_total']} pontos, dos quais {R['q_outliers_yoy_em_pandemia_2020_21']} em mar/2020–jun/2021 (choque e rebote: maior valor, veículos em abr/2021, +{f(R['q_outliers_yoy_top'][0]['var_yoy_pct'], 0)}% sobre a base do lockdown) e {R['q_outliers_yoy_fora_pandemia']} fora dele.",
     "A janela de pandemia é excluída das estatísticas (P2 a P5); os 10 pontos restantes são mantidos (variações reais de setores voláteis)."],
    ["Atualidade", "Defasagem entre fontes", f"PIM-PF vai até jul/2026; o BCB, até {R['q_macro_ultimo_mes'][:7].replace('-', '/')} (defasagem de {R['q_defasagem_pim_vs_macro_meses']} meses). O último mês do BCB é parcial.",
     "Junção por interseção (R9); o cenário P5 usa o último juro real completo (ago./2026)."],
])
im(EVID / "fig6_qualidade_ausencias.png", "Figura 1 — Meses sem informação no IBGE: apenas as divisões 3.18 e 3.33, antes de 2012.")
p("Ressalva importante: o índice não tem ajuste sazonal. A série industrial oscila fortemente entre meses por sazonalidade (Figura 2), o que torna inadequada a comparação mês a mês. "
  "Por isso toda a análise usa variação contra o mesmo mês do ano anterior, que elimina a sazonalidade, e médias de 12 meses.")
im(EVID / "fig1_industria_e_juro_real.png", "Figura 2 — Produção industrial (índice sem ajuste sazonal) e juro real ex-post, 2002–2026.")
p(f"Amostra de análise. Para P2 a P5 usa-se janeiro/2013 a julho/2026, excluindo mar/2020–jun/2021: {R['amostra_meses']} meses, comum às 27 atividades. "
  "O corte em 2013 é o primeiro mês em que as duas divisões de série curta têm variação anual; a exclusão da pandemia evita que um único choque de +900% domine correlações e regressões.")

h2("5.2 P1 — Recuperação em relação a 2019")
p(f"O agregado voltou: a indústria geral produz hoje (média dos últimos 12 meses) {f((P1r['geral_razao'] - 1) * 100)}% acima da média de 2019, com queda máxima de "
  f"{f(abs(P1r['geral_queda_2020']), 0)}% no fundo de 2020 e retorno sustentado ao patamar de 2019 em {f(P1r['geral_meses'], 0)} meses. As extrativas estão "
  f"{f((P1r['extrativas'] - 1) * 100)}% acima, e a indústria de transformação, {f((P1r['transformacao'] - 1) * 100)}% acima. "
  f"A média, porém, esconde a divisão: das {P1r['n_divisoes']} divisões da transformação, {P1r['abaixo']} ainda estão abaixo de 2019 e {P1r['acima']} acima (Figura 3).")
im(EVID / "fig2_P1_recuperacao_pre_pandemia.png", "Figura 3 — Produção atual vs. média de 2019 por divisão (P1).")
sel = P1[P1.codigo.str.startswith("3.")].sort_values("razao_ult12m_2019")
linhas = pd.concat([sel.head(5), sel.tail(5)])
tb(["Cód.", "Divisão", "Atual vs. 2019", "Queda máx. 2020", "Meses p/ recuperar"],
   [[r.codigo, r.atividade, f"{(r.razao_ult12m_2019 - 1) * 100:+.1f}%".replace(".", ","), f"{r.queda_maxima_2020_pct:.0f}%".replace(".", ","),
     "não recuperou" if pd.isna(r.meses_ate_recuperar) else f"{int(r.meses_ate_recuperar)}"] for r in linhas.itertuples()])
p("Resposta a P1. Os cinco piores casos são Impressão (−33%), Confecção (−27%), Móveis (−17%), Couro e calçados (−16%) e Produtos diversos (−16%): atividades intensivas em mão de obra e "
  "ligadas ao consumo das famílias e à impressão de mídia física. Os cinco melhores são Fumo (+38%), Derivados de petróleo (+15%), Máquinas e equipamentos (+11%), "
  "Outros equipamentos de transporte (+10%) e Celulose e papel (+8%), setores ligados a commodities, energia e bens de capital. "
  "Quatro divisões (Impressão, Confecção, Produtos diversos e Manutenção) nunca tiveram três meses seguidos acima da média de 2019 desde o choque. "
  "Veículos automotores levou 52 meses para recuperar o patamar, e Outros equipamentos de transporte, 29, embora ambos tenham a maior queda inicial (−92% e −81%). "
  "H3 confirmada: a recuperação foi desigual, e o agregado (+3,0%) esconde 13 divisões abaixo do nível pré-pandemia.")
p("Cautela: comparar com a média de um único ano (2019) é uma escolha de base. 2019 foi um ano de crescimento fraco, então \"acima de 2019\" não significa \"acima da tendência\".")

h2("5.3 P2 — Sensibilidade ao juro real")
p("Para cada uma das 27 atividades calculou-se a correlação entre o crescimento anual em t e o juro real em t−k, para k de 0 a 12 meses (Figura 4). O sinal esperado por H1 é negativo. "
  "Como a escolha do melhor k dentre 13 defasagens infla a chance de achar correlação por acaso, e as séries são altamente autocorrelacionadas, a significância foi avaliada por um teste de permutação: "
  f"o juro real é deslocado circularmente ({P2r['perm_n']:,} deslocamentos aleatórios) e compara-se a correlação mais negativa observada com a distribuição do mínimo sobre as 13 defasagens.".replace(",", "."))
im(EVID / "fig3_P2_correlacao_por_defasagem.png", "Figura 4 — Correlação juro real × crescimento anual por defasagem, divisões da transformação (P2).")
top = P2r["top"]
tb(["Atividade", "Defasagem", "Correlação", "p (permutação)"],
   [[f"{c} {n}", str(k), f"{r:+.2f}".replace(".", ","), f(pv, 2)] for c, n, r, k, pv in top] +
   [["1 Indústria geral", str(P2r["geral"]["lag"]), f"{P2r['geral']['r']:+.2f}".replace(".", ","), f(P2r["geral"]["p"], 2)]])
p(f"Resposta a P2. Não há evidência de que o juro real explique o crescimento anual em nenhuma atividade: {P2r['n_significativas']} das {P2r['n_total']} séries têm p < 0,05 no teste de permutação. "
  f"Só {P2r['n_negativas']} das {P2r['n_total']} correlações têm o sinal esperado, e a mais forte (Produtos químicos, defasagem de 2 meses) é de apenas −0,21, com p = 0,21. "
  f"Para a indústria geral a correlação é positiva (+{f(P2r['geral']['r'], 2)}), o oposto de H1. A defasagem mediana ótima é {f(P2r['lag_mediano_divisoes'], 0)} mês, ou seja, não há padrão temporal identificável. "
  "H1 não é sustentada por estes dados.")
p("Interpretação. Três explicações são plausíveis, e a análise não separa entre elas: (i) endogeneidade, porque o Banco Central sobe juros em ciclos de demanda aquecida e corta em recessões, o que gera correlação positiva espúria; "
  "(ii) o efeito do crédito passa por canais que o juro real ex-post não captura (spread, acesso a crédito, expectativas); (iii) amostra curta, de 13 anos, com poucos ciclos de juros. "
  "O resultado negativo é informativo: quem usar a Selic corrente para prever a produção setorial no curto prazo não terá respaldo nestes dados.")

h2("5.4 P3 — Câmbio, controlando o juro")
p("Para cada atividade estimou-se por mínimos quadrados o crescimento anual contra o juro real defasado (no lag ótimo de P2) e a variação cambial de 12 meses, com erros-padrão robustos a autocorrelação e heterocedasticidade (Newey-West, 12 defasagens). A Figura 5 mostra os coeficientes com intervalos de 95%.")
im(EVID / "fig4_P3_coeficientes.png", "Figura 5 — Coeficientes de juro real e de câmbio por divisão (P3).")
cn = P3r["cambio_neg_sig"]
tb(["Atividade", "β câmbio (p.p. por 1% de desvalorização)", "estatística t"],
   [[f"{c} {n}", f"{b:+.2f}".replace(".", ","), f"{t:+.2f}".replace(".", ",")] for c, n, b, t in cn] +
   [["1 Indústria geral", f"{P3r['geral']['beta_cambio']:+.2f}".replace(".", ","), f"{P3r['geral']['t_cambio']:+.2f}".replace(".", ",")]])
p(f"Resposta a P3. O câmbio tem associação estatisticamente significativa em {P3r['n_cambio_sig']} das {P3r['n_total']} atividades, e em todas o sinal é negativo: desvalorização do real acompanha crescimento anual menor. "
  f"Para a indústria geral, cada 1% de desvalorização em 12 meses está associado a {f(abs(P3r['geral']['beta_cambio']), 2)} p.p. a menos de crescimento anual (t = {f(P3r['geral']['t_cambio'], 1)}). "
  "Os maiores efeitos estão em Informática/eletrônicos, Veículos e Impressão, que dependem de insumos e componentes importados. Não há atividade com efeito positivo significativo, portanto a parte de H2 sobre exportadores não se confirma. "
  f"A ressalva: o poder explicativo é baixo (R² mediano {f(P3r['r2_mediano'], 2)}; máximo {f(P3r['r2_max'], 2)}, em {P3r['r2_max_nome'].strip()}), e as grandes desvalorizações do período (2015 e 2020) coincidem com recessões, "
  "de modo que o coeficiente pode refletir o ciclo econômico em vez do câmbio em si.")

h2("5.5 P4 — Volatilidade")
sd = P4r
p("A volatilidade foi medida pelo desvio-padrão do crescimento anual na amostra de análise. As divisões variam mais de quatro vezes entre si: "
  f"Fumo ({f(sd['mais_volateis'][0][2], 1)} p.p.), Informática ({f(sd['mais_volateis'][1][2], 1)}), Veículos ({f(sd['mais_volateis'][2][2], 1)}) e Impressão ({f(sd['mais_volateis'][3][2], 1)}) "
  f"são as mais voláteis; Celulose e papel ({f(sd['menos_volateis'][0][2], 1)}), Químicos ({f(sd['menos_volateis'][1][2], 1)}) e Minerais não metálicos ({f(sd['menos_volateis'][2][2], 1)}) são as mais estáveis. "
  f"A indústria geral, por diversificação, tem {f(sd['geral'], 1)} p.p., no nível da divisão mais estável.")
p(f"Resposta a P4. A razão entre a divisão mais e a menos volátil é de {f(sd['razao_max_min'], 1)} vezes. O pior crescimento anual de Veículos é −39% sem a pandemia e −92% com ela; o de Celulose e papel é −11% nos dois casos. "
  f"A volatilidade tem correlação fraca com a sensibilidade ao juro ({f(sd['corr_beta_vol'], 2)}), então o risco setorial vem sobretudo de fatores próprios da atividade, não da política monetária. "
  "Em termos de decisão, esta é a conclusão mais robusta do trabalho: o desvio-padrão setorial é estável e diretamente mensurável, ao contrário da sensibilidade ao juro.")

h2("5.6 P5 — Cenário de queda do juro real e discussão geral")
p(f"O cenário aplica ao coeficiente de juro de cada atividade uma queda de 3 p.p. (de {f(R['P5_juro_real_atual_pct'], 1)}% a.a. em {R['P5_ult_mes_juro_real'][:7].replace('-', '/')}, com Selic de {f(R['P5_selic_atual'], 2)}% "
  f"e IPCA 12m de {f(R['P5_ipca12m_atual'], 2)}%, para cerca de {f(R['P5_juro_real_atual_pct'] - 3, 1)}%, próximo à média histórica de {f(R['P5_juro_real_media_hist_pct'], 1)}%). O intervalo de 95% usa o erro-padrão robusto.")
im(EVID / "fig5_P4_P5_sensibilidade_volatilidade.png", "Figura 6 — Efeito projetado de −3 p.p. de juro real vs. volatilidade (P4/P5).")
sel5 = P5[P5.codigo.isin(["1", "3", "3.28", "3.19", "3.18", "3.31"])]
ordem = ["1", "3", "3.28", "3.19", "3.18", "3.31"]
sel5 = sel5.set_index("codigo").loc[ordem].reset_index()
tb(["Atividade", "Efeito projetado (p.p.)", "IC 95%", "IC exclui zero?"],
   [[f"{r.codigo} {r.atividade}", f"{r.efeito_pp:+.2f}".replace(".", ","), f"[{r.ic_inf:.2f}; {r.ic_sup:.2f}]".replace(".", ","), "sim" if r.ic_exclui_zero else "não"] for r in sel5.itertuples()])
p(f"Resposta a P5. Para a indústria geral o modelo projeta {f(P5r['geral']['efeito'], 2)} p.p. de crescimento anual (IC 95%: {f(P5r['geral']['ic'][0], 1)} a +{f(P5r['geral']['ic'][1], 1)}), isto é, "
  f"efeito indistinguível de zero e com sinal contrário ao esperado. Entre as {P5r['n_div']} divisões, {P5r['n_efeito_positivo']} têm efeito projetado positivo, mas em apenas {P5r['n_ic_exclui_zero']} o intervalo exclui zero: "
  "Máquinas e equipamentos (+2,4 p.p.; IC +0,3 a +4,6; defasagem de 2 meses) e Derivados de petróleo (+1,7 p.p.; IC +0,01 a +3,3). Com 24 testes a 5%, esperam-se cerca de 1,2 falsos positivos por acaso, "
  "e nenhuma das duas sobrevive à correção do P2 (p = 0,21 e 0,27). Os efeitos de maior magnitude (Impressão +3,2 p.p.) têm IC que inclui grandes valores negativos. "
  "As extrativas têm efeito significativo com sinal invertido (−1,2 p.p.), consistente com o ciclo de commodities, e não com causalidade dos juros.")
p("Discussão geral. Os resultados formam um quadro coerente, embora em parte negativo. (1) O diagnóstico descritivo é sólido e útil: o agregado industrial voltou a 2019, mas 13 de 24 divisões não, com um núcleo de "
  "setores de consumo e mídia ainda 16% a 33% abaixo. (2) A hipótese de que juros ditam a produção setorial não se sustentou: nenhuma série apresenta relação distinguível do acaso, e o agregado tem sinal oposto ao esperado, "
  "o que é compatível com endogeneidade da política monetária. (3) O câmbio tem associação negativa e consistente, mas de baixo poder explicativo. (4) A volatilidade setorial, que varia mais de quatro vezes, é a "
  "informação mais estável para decisão. Para um gestor, a implicação prática é que priorizar setores por sensibilidade ao corte de juros não tem base empírica neste modelo; "
  "Máquinas e equipamentos e Derivados de petróleo são hipóteses a investigar com dados adicionais (crédito, investimento, preços), não recomendações. O que os dados sustentam é a leitura de risco (P4) e a de recuperação (P1).")

# =====================================================================================================
h1("6. Ética, licenciamento e proteção de dados")
bl([
    "Licença dos dados. Ambas as fontes são estatísticas públicas divulgadas como dados abertos: IBGE (SIDRA) e Banco Central do Brasil (SGS). Exigem citação da fonte. A licença específica de cada conjunto não foi verificada automaticamente nesta execução, e deve ser conferida em sidra.ibge.gov.br e dadosabertos.bcb.gov.br antes de qualquer reutilização comercial. Por serem séries agregadas de tamanho pequeno (cerca de 1,7 MB), os arquivos brutos foram versionados para garantir reprodutibilidade; o script de coleta também está publicado.",
    "Dados pessoais (LGPD). Não se aplica: as bases são agregados nacionais mensais, sem identificadores de pessoas ou empresas.",
    "Dados corporativos. Não se aplica: nenhum dado de empresa foi utilizado.",
    "Coleta automatizada. Foram usadas apenas APIs oficiais, sem scraping, com pausa entre tentativas.",
    "Código. O código do repositório está sob a licença MIT (arquivo LICENSE).",
])

# =====================================================================================================
h1("7. Autoavaliação")
p("Objetivos atingidos. As perguntas P1, P3 e P4 foram respondidas com evidência. P2 e P5 foram respondidas, mas com resultado negativo/inconclusivo: os dados não sustentam relação "
  "entre juro real e produção setorial. O resultado é legítimo e foi reportado como tal, sem remoção de perguntas nem busca de especificações até achar significância.")
p("Perguntas sugeridas:")
bl([
    "Quais perguntas foram respondidas e quais não? P1, P3 e P4 foram respondidas de forma conclusiva. P2 e P5 tiveram resposta negativa: não há relação detectável entre juro real e crescimento setorial. A causa provável é a combinação de endogeneidade da política monetária, amostra curta e um regressor (juro real ex-post) que não capta o canal de crédito.",
    "Que limitações dos dados condicionaram os resultados? Índice sem ajuste sazonal (obrigou o uso de variação anual, que perde informação de curto prazo); apenas 13 anos comparáveis; efeito da pandemia, que exigiu excluir 16 meses; ausência de controles (crédito, renda, demanda externa, preços de commodities); e nível nacional agregado, sem heterogeneidade regional. O câmbio e o juro são medidos por séries únicas para todas as atividades, embora a exposição varie por setor (exportador, importador).",
    "Que decisões técnicas tomaria de outra forma? Testaria o pipeline no Databricks desde o primeiro dia, em vez de validá-lo localmente antes. Só na execução real apareceu que o Free Edition não acessa a internet, o que obrigou a separar a coleta (local) da persistência (nuvem) e adaptar o notebook 01; descobrir isso cedo teria orientado o desenho da coleta. Usaria a série com ajuste sazonal que o IBGE também divulga e um modelo em primeiras diferenças ou VAR, mais adequado a séries persistentes. Incluiria ao menos uma variável de crédito (concessões do BCB) e uma de demanda externa.",
    "Que extensões transformariam o MVP em solução de uso contínuo? Agendamento mensal do job (Databricks Workflows) alinhado à divulgação da PIM-PF; carga incremental em vez de recarga completa; testes automatizados de qualidade (as verificações da Seção 5.1 viram expectativas de dados); painel de monitoramento; inclusão de PIM regional e de indicadores de crédito; e monitoramento da estabilidade das estimativas ao longo do tempo.",
])
p("Trabalhos futuros. Estimar o efeito do juro com identificação mais forte (choques monetários de alta frequência, ou variáveis instrumentais); incluir a pesquisa regional (PIM-PF regional) e a de comércio exterior para separar exportadores de importadores no P3; estender a análise a serviços e comércio.")

h1("8. Reprodutibilidade")
bl([
    "Local: python scripts/coleta.py → python scripts/etl.py → python scripts/analise.py → python scripts/gerar_catalogo.py → python scripts/gerar_relatorio.py. Dependências em requirements.txt. Semente aleatória fixa (42).",
    "Databricks: criar uma Git folder com o repositório e executar notebooks/01, 02 e 03 em ordem. Os notebooks usam catalog workspace e schema industria_juros. O Free Edition não tem acesso à internet: o notebook 01 tenta a API e, se falhar, usa o bronze versionado em data/raw.",
    f"Repositório: {REPO}",
    "Fontes: IBGE/SIDRA tabela 8888 (PIM-PF); BCB/SGS séries 4189, 433, 13522 e 3698.",
])


# =====================================================================================================
# Renderização
def md_render():
    out = ["# MVP — Pipeline de dados em nuvem: produção industrial × juros, câmbio e inflação\n",
           "**Universidade de Brasília — Departamento de Engenharia de Produção · Seminários (SSD/MVP)**  \n"
           "Aluno: Lucas Iwano · Professor: André Luiz Marques Serrano\n",
           f"> Repositório: {REPO}  \n> Relatório em Word: `Relatorio_MVP_Industria_Juros_Lucas_Iwano.docx`\n",
           "## Como reproduzir (resumo)\n",
           "```bash\npip install -r requirements.txt\npython scripts/coleta.py && python scripts/etl.py && python scripts/analise.py\n```\n",
           "No Databricks: crie uma *Git folder* com este repositório e execute `notebooks/01`, `02` e `03` em ordem.\n"]
    for b in B:
        t = b[0]
        if t == "h1":
            out.append(f"\n## {b[1]}\n")
        elif t == "h2":
            out.append(f"\n### {b[1]}\n")
        elif t == "p":
            out.append(b[1] + "\n")
        elif t == "ul":
            out.append("\n".join(f"- {i}" for i in b[1]) + "\n")
        elif t == "table":
            head, rows = b[1], b[2]
            out.append("| " + " | ".join(head) + " |\n|" + "---|" * len(head))
            out.extend("| " + " | ".join(str(c).replace("|", "/") for c in r) + " |" for r in rows)
            out.append("")
        elif t == "img":
            rel = Path(b[1]).relative_to(ROOT).as_posix()
            out.append(f"![{b[2]}]({rel})\n\n*{b[2]}*\n")
    return "\n".join(out)


def shade(cell, cor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), cor)
    tcPr.append(sh)


def docx_render():
    d = Document()
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.12
    for nm, sz, cor in [("Heading 1", 16, "1F4E79"), ("Heading 2", 13, "1F4E79")]:
        s = d.styles[nm]
        s.font.name = "Calibri"
        s.font.size = Pt(sz)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(cor)
    # capa
    for txt, sz, bold in [("UNIVERSIDADE DE BRASÍLIA", 14, True), ("Departamento de Engenharia de Produção", 12, False), ("", 11, False),
                          ("MVP — MINIMUM VIABLE PRODUCT", 20, True),
                          ("Pipeline de dados em nuvem: produção industrial brasileira, juros, câmbio e inflação", 16, True),
                          ("Recuperação pós-pandemia, sensibilidade ao juro real e risco por setor da indústria", 12, False), ("", 11, False),
                          ("Aluno: Lucas Iwano", 12, False), ("Professor: André Luiz Marques Serrano", 12, False),
                          ("Setembro de 2026", 12, False), (f"Repositório: {REPO}", 11, False)]:
        par = d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run(txt)
        r.bold = bold
        r.font.size = Pt(sz)
        if sz >= 16:
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    d.add_page_break()
    for b in B:
        t = b[0]
        if t == "h1":
            d.add_heading(b[1], 1)
        elif t == "h2":
            d.add_heading(b[1], 2)
        elif t == "p":
            par = d.add_paragraph()
            txt = b[1]
            partes = txt.split("**")
            for i, seg in enumerate(partes):
                run = par.add_run(seg)
                run.bold = i % 2 == 1
        elif t == "ul":
            for i in b[1]:
                d.add_paragraph(i, style="List Bullet")
        elif t == "table":
            head, rows = b[1], b[2]
            tab = d.add_table(rows=1, cols=len(head))
            tab.style = "Table Grid"
            tab.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, hd in enumerate(head):
                c = tab.rows[0].cells[i]
                c.text = ""
                run = c.paragraphs[0].add_run(hd)
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(255, 255, 255)
                shade(c, "1F4E79")
            for r in rows:
                cells = tab.add_row().cells
                for i, v in enumerate(r):
                    cells[i].text = ""
                    run = cells[i].paragraphs[0].add_run(str(v))
                    run.font.size = Pt(8.5)
            pesos = [min(max(len(str(x)) for x in [head[i]] + [r[i] for r in rows]), 38) + 12 for i in range(len(head))]
            tab.autofit = False
            for i, w in enumerate(pesos):
                for row in tab.rows:
                    row.cells[i].width = Cm(16.6 * w / sum(pesos))
            d.add_paragraph()
        elif t == "img":
            d.add_picture(str(b[1]), width=Cm(15.5))
            d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = d.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = cap.add_run(b[2])
            rr.italic = True
            rr.font.size = Pt(9)
    d.save(DOCX_OUT)


(ROOT / "README.md").write_text(md_render(), encoding="utf-8")
docx_render()
print("README.md e", DOCX_OUT.name, "gerados;", len(B), "blocos")
